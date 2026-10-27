from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from conference_tools import linkedin
from conference_tools.linkedin import import_linkedin_profiles, parse_json_object
from conference_tools.workbook import SHEET_SPEAKERS


@pytest.fixture
def blank_profile(template_path):
    """Classeur où Ada n'a ni poste, ni bio, ni photo."""
    workbook = load_workbook(template_path)
    ws = workbook[SHEET_SPEAKERS]
    for column in ("G", "H", "I", "J", "K"):
        ws[f"{column}2"] = None
    workbook.save(template_path)
    return template_path


@pytest.fixture
def pdf_dir(tmp_path, monkeypatch) -> Path:
    directory = tmp_path / "pdf"
    directory.mkdir()
    (directory / "Ada_Lovelace.pdf").write_bytes(b"%PDF")
    monkeypatch.setattr(linkedin, "extract_pdf_text", lambda path: f"profil {path.name}")
    return directory


def fake_extractor(text: str) -> dict:
    assert "Ada_Lovelace.pdf" in text
    return {
        "jobTitle": "Chercheuse",
        "company": "Analytical Engines",
        "bio": "Ada Lovelace est chercheuse.",
        "expertise": ["IA", " ", "Maths"],
    }


def speaker_row(path: Path, row: int = 2) -> dict[str, object]:
    ws = load_workbook(path)[SHEET_SPEAKERS]
    headers = [cell.value for cell in ws[1]]
    return dict(zip(headers, [cell.value for cell in ws[row]], strict=True))


def test_fills_empty_cells_from_pdf_and_photos(blank_profile, pdf_dir, tmp_path):
    photos = tmp_path / "photos"
    photos.mkdir()
    (photos / "ada-lovelace.webp").write_bytes(b"img")

    report = import_linkedin_profiles(blank_profile, pdf_dir, fake_extractor, photos_dir=photos)

    assert report.updated == ["Ada Lovelace"]
    assert report.photos == ["Ada Lovelace"]
    assert "Alan Turing : fiche déjà complète" in report.skipped
    row = speaker_row(blank_profile)
    assert row["Poste"] == "Chercheuse"
    assert row["Expertises"] == "IA; Maths"
    assert row["URL photo"] == "/images/speakers/ada-lovelace.webp"


def test_does_not_overwrite_without_flag(template_path, pdf_dir):
    report = import_linkedin_profiles(template_path, pdf_dir, fake_extractor)
    assert any("déjà complète" in s for s in report.skipped)
    assert speaker_row(template_path)["Poste"] == "Directrice scientifique"

    import_linkedin_profiles(template_path, pdf_dir, fake_extractor, overwrite=True)
    assert speaker_row(template_path)["Poste"] == "Chercheuse"


def test_skips_speakers_without_consent(blank_profile, pdf_dir):
    workbook = load_workbook(blank_profile)
    workbook[SHEET_SPEAKERS]["E2"] = "Non"
    workbook.save(blank_profile)
    report = import_linkedin_profiles(blank_profile, pdf_dir, fake_extractor)
    assert "Ada Lovelace : pas de consentement" in report.skipped
    assert speaker_row(blank_profile)["Poste"] is None


def test_dry_run_and_extraction_errors(blank_profile, pdf_dir):
    def failing(_: str) -> dict:
        raise ValueError("JSON invalide")

    report = import_linkedin_profiles(blank_profile, pdf_dir, failing, dry_run=True)
    assert any("échec de l'extraction (JSON invalide)" in s for s in report.skipped)

    import_linkedin_profiles(blank_profile, pdf_dir, fake_extractor, dry_run=True)
    assert speaker_row(blank_profile)["Poste"] is None


def test_parse_json_object():
    assert parse_json_object('Voici :\n```json\n{"a": 1}\n```') == {"a": 1}
    with pytest.raises(ValueError):
        parse_json_object("pas de json")
