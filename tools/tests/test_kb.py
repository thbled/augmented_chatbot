from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

from conference_tools.kb import build_knowledge_base, format_day
from conference_tools.models import truncate
from conference_tools.workbook import SHEET_SPEAKERS, read_workbook

CARD_RE = re.compile(r"```(speaker|session|partner)\n(.+?)\n```", re.DOTALL)


def cards_in(path: Path) -> list[tuple[str, dict]]:
    return [(kind, json.loads(body)) for kind, body in CARD_RE.findall(path.read_text())]


def test_build_knowledge_base(template_path, tmp_path):
    data, _ = read_workbook(template_path)
    static = tmp_path / "static"
    (static / "documentation").mkdir(parents=True)
    (static / "assistant.md").write_text("# Assistant")
    (static / "documentation" / "rag.md").write_text("# RAG")
    (static / "documentation" / "README.md").write_text("ignoré")

    out = tmp_path / "kb"
    (out / "obsolete.md").parent.mkdir()
    (out / "obsolete.md").write_text("ancien")
    files = build_knowledge_base(data, out, static)

    names = {str(path.relative_to(out)) for path in files}
    assert {
        "evenement.md",
        "programme.md",
        "speakers.md",
        "partenaires.md",
        "infos_pratiques.md",
        "speakers/speaker_ada-lovelace.md",
        "partenaires/partenaire_acme-cloud.md",
        "assistant.md",
        "documentation/rag.md",
    } <= names
    assert not (out / "obsolete.md").exists()
    assert not (out / "documentation" / "README.md").exists()
    assert len([n for n in names if n.startswith("sessions/")]) == 3


def test_speaker_document_contains_card(template_path, tmp_path):
    data, _ = read_workbook(template_path)
    build_knowledge_base(data, tmp_path / "kb")
    path = tmp_path / "kb" / "speakers" / "speaker_ada-lovelace.md"
    text = path.read_text()
    assert "# Speaker : Ada Lovelace — Directrice scientifique, Analytical Engines" in text
    assert "jeudi 12 novembre 2026, 09:30–10:15" in text
    assert "avec Alan Turing" in text

    [(kind, card)] = cards_in(path)
    assert kind == "speaker"
    assert card["id"] == "ada-lovelace"
    assert card["photoUrl"] == "/images/speakers/ada-lovelace.jpg"
    assert [s["start"] for s in card["sessions"]] == ["09:30", "10:30"]
    assert "consent" not in json.dumps(card)


def test_session_card_and_unpublished_speakers(template_path, tmp_path):
    workbook = load_workbook(template_path)
    workbook[SHEET_SPEAKERS]["E3"] = "Non"  # Alan Turing retire son consentement
    workbook.save(template_path)
    data, _ = read_workbook(template_path)
    build_knowledge_base(data, tmp_path / "kb")

    assert not (tmp_path / "kb" / "speakers" / "speaker_alan-turing.md").exists()
    session_file = next((tmp_path / "kb" / "sessions").glob("*securiser*"))
    [(kind, card)] = cards_in(session_file)
    assert kind == "session"
    # Le speaker non publié n'apparaît que par son nom, sans identifiant ni photo.
    assert card["speakers"] == [
        {"name": "Alan Turing"},
        {
            "id": "ada-lovelace",
            "name": "Ada Lovelace",
            "photoUrl": "/images/speakers/ada-lovelace.jpg",
        },
    ]
    assert card["date"] == "2026-11-12"


def test_program_overview_is_a_table(template_path, tmp_path):
    data, _ = read_workbook(template_path)
    build_knowledge_base(data, tmp_path / "kb")
    text = (tmp_path / "kb" / "programme.md").read_text()
    assert "## Programme du jeudi 12 novembre 2026" in text
    assert "| 09:30–10:15 | L'IA générative en entreprise" in text
    assert "### Cybersécurité" in text


def test_truncate_and_format_day():
    assert truncate("un deux trois quatre", 12) == "un deux…"
    assert truncate("court", 12) == "court"
    from datetime import date

    assert format_day(date(2026, 11, 12)) == "jeudi 12 novembre 2026"


def test_build_scope(template_path, tmp_path):
    from conference_tools.kb import build_scope, write_scope

    workbook = load_workbook(template_path)
    workbook[SHEET_SPEAKERS]["E3"] = "Non"  # Alan Turing : nom conservé (programme public)
    workbook.save(template_path)
    data, _ = read_workbook(template_path)
    scope = build_scope(data)
    assert scope["eventName"] == "Business Agility Days 2026"
    assert scope["themes"] == ["IA générative", "Data", "Cloud", "Cybersécurité"]
    assert scope["speakers"] == ["Ada Lovelace", "Alan Turing"]
    assert scope["partners"] == ["Acme Cloud"]
    assert "Sécuriser ses applications LLM" in scope["sessions"]
    path = write_scope(data, tmp_path / "out" / "scope.json")
    assert json.loads(path.read_text())["partners"] == ["Acme Cloud"]
