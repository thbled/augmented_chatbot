from __future__ import annotations

from datetime import date, time

import pytest
from openpyxl import load_workbook

from conference_tools.workbook import (
    SHEET_PROGRAM,
    SHEET_SPEAKERS,
    WorkbookError,
    create_template,
    parse_date,
    parse_time,
    read_workbook,
)


def test_template_contains_all_sheets(template_path):
    workbook = load_workbook(template_path)
    assert workbook.sheetnames == [
        "Instructions",
        "Événement",
        "Speakers",
        "Programme",
        "Partenaires",
        "Infos pratiques",
    ]
    headers = [cell.value for cell in workbook[SHEET_SPEAKERS][1]]
    assert "Prénom *" in headers
    assert "Consentement *" in headers
    assert workbook[SHEET_SPEAKERS]["B1"].comment is not None


def test_read_example_template(template_path):
    data, report = read_workbook(template_path)
    assert report.ok, report.errors
    assert data.event.name == "Business Agility Days 2026"
    assert [s.id for s in data.speakers] == ["ada-lovelace", "alan-turing"]
    assert data.speakers[0].expertise == ["IA générative", "Algorithmique", "Gouvernance de l'IA"]
    assert len(data.sessions) == 3
    assert data.sessions[1].start == time(9, 30)
    # Ada intervient dans deux sessions, triées chronologiquement.
    assert [s.start for s in data.speakers[0].sessions] == [time(9, 30), time(10, 30)]
    assert data.partners[0].id == "acme-cloud"
    assert len(data.practical) == 2


def test_empty_template_reads_without_data(tmp_path):
    path = tmp_path / "vide.xlsx"
    create_template(path, with_examples=False)
    data, report = read_workbook(path)
    assert report.ok
    assert data.speakers == data.sessions == data.partners == []
    assert any("nom de l'événement" in w for w in report.warnings)


def test_reports_errors_and_warnings(template_path):
    workbook = load_workbook(template_path)
    speakers = workbook[SHEET_SPEAKERS]
    speakers["B4"] = "Grace"  # prénom sans nom → erreur
    speakers["B5"], speakers["C5"], speakers["E5"] = "Sans", "Accord", "Non"
    speakers["G5"] = "ftp://photo"
    program = workbook[SHEET_PROGRAM]
    program["B5"], program["C5"], program["D5"], program["E5"] = "12/11/2026", "11:00", "10:00", "X"
    program["B6"], program["C6"], program["D6"] = "12/11/2026", "14h", "15h"
    program["E6"], program["I6"] = "Atelier", "Inconnu Total"
    workbook.save(template_path)

    data, report = read_workbook(template_path)
    assert any("Speakers ligne 4" in e for e in report.errors)
    assert any("Programme ligne 5" in e and "fin" in e for e in report.errors)
    assert any("pas de consentement" in w for w in report.warnings)
    assert any("URL ignorée" in w for w in report.warnings)
    assert any("Inconnu Total" in w for w in report.warnings)
    atelier = next(s for s in data.sessions if s.title == "Atelier")
    assert (atelier.start, atelier.end) == (time(14, 0), time(15, 0))


def test_homonyms_get_distinct_ids(template_path):
    workbook = load_workbook(template_path)
    speakers = workbook[SHEET_SPEAKERS]
    speakers["B4"], speakers["C4"], speakers["E4"] = "Ada", "Lovelace", "Oui"
    workbook.save(template_path)
    data, _ = read_workbook(template_path)
    assert [s.id for s in data.speakers] == ["ada-lovelace", "alan-turing", "ada-lovelace-2"]


def test_missing_sheet_raises(tmp_path):
    path = tmp_path / "incomplet.xlsx"
    create_template(path)
    workbook = load_workbook(path)
    del workbook[SHEET_PROGRAM]
    workbook.save(path)
    with pytest.raises(WorkbookError, match="Programme"):
        read_workbook(path)


@pytest.mark.parametrize(
    ("value", "expected"),
    [("12/11/2026", date(2026, 11, 12)), ("2026-11-12", date(2026, 11, 12)), ("demain", None)],
)
def test_parse_date(value, expected):
    assert parse_date(value) == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [("9:30", time(9, 30)), ("14h", time(14, 0)), ("14h05", time(14, 5)), ("25:00", None)],
)
def test_parse_time(value, expected):
    assert parse_time(value) == expected


def test_create_template_creates_parent_dirs(tmp_path):
    path = tmp_path / "a" / "b" / "c.xlsx"
    create_template(path)
    assert path.exists()
