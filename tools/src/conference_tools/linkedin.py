"""Pré-remplissage des fiches speakers à partir des exports PDF de profils LinkedIn.

Aucun scraping : on part du PDF généré par LinkedIn (« Plus » > « Enregistrer au format PDF »),
fourni par le speaker ou téléchargé par l'organisateur avec son autorisation. Le texte est
structuré par le LLM (API compatible OpenAI), puis écrit dans les cellules vides du classeur.
Seuls les speakers dont le consentement vaut « Oui » sont traités.
"""

from __future__ import annotations

import json
import re
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

from conference_tools.models import slugify
from conference_tools.workbook import SHEET_SPEAKERS, SPEAKER_COLUMNS, cell_text, is_yes

PHOTO_EXTENSIONS = (".jpg", ".jpeg", ".png", ".webp", ".avif")
MAX_PROFILE_CHARS = 12000

EXTRACTION_PROMPT = """Tu prépares la fiche d'un speaker pour une conférence professionnelle.
À partir du texte d'un export PDF de profil LinkedIn, renvoie UNIQUEMENT un objet JSON :
{
  "jobTitle": "intitulé du poste actuel",
  "company": "entreprise actuelle",
  "bio": "biographie professionnelle en français, à la 3e personne, 400 à 800 caractères",
  "expertise": ["3 à 6 domaines d'expertise courts"]
}
Règles : n'utilise que les informations du profil, n'invente rien, aucune donnée personnelle
(adresse, téléphone, e-mail, âge, situation familiale). Champ inconnu : chaîne vide ou liste vide.
"""

FIELD_MAP = {
    "job_title": "jobTitle",
    "company": "company",
    "bio": "bio",
    "expertise": "expertise",
}

Extractor = Callable[[str], dict[str, Any]]


@dataclass
class ImportReport:
    updated: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)
    photos: list[str] = field(default_factory=list)


def extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    return text[:MAX_PROFILE_CHARS]


def parse_json_object(raw: str) -> dict[str, Any]:
    match = re.search(r"\{.*\}", raw, re.DOTALL)
    if not match:
        raise ValueError("Aucun objet JSON dans la réponse du LLM")
    value = json.loads(match.group(0))
    if not isinstance(value, dict):
        raise ValueError("La réponse du LLM n'est pas un objet JSON")
    return value


def llm_extractor(base_url: str, api_key: str, model: str) -> Extractor:
    from openai import OpenAI

    client = OpenAI(base_url=base_url, api_key=api_key, timeout=120)

    def extract(profile_text: str) -> dict[str, Any]:
        response = client.chat.completions.create(
            model=model,
            temperature=0.2,
            messages=[
                {"role": "system", "content": EXTRACTION_PROMPT},
                {"role": "user", "content": profile_text},
            ],
        )
        return parse_json_object(response.choices[0].message.content or "")

    return extract


def find_file(directory: Path, speaker_id: str, extensions: tuple[str, ...]) -> Path | None:
    """Cherche `<id>.<ext>`, ou à défaut un fichier dont le nom contient l'identifiant."""
    if not directory.is_dir():
        return None
    candidates = [p for p in directory.iterdir() if p.suffix.lower() in extensions]
    for path in candidates:
        if slugify(path.stem) == speaker_id:
            return path
    for path in candidates:
        if speaker_id in slugify(path.stem):
            return path
    return None


def import_linkedin_profiles(
    workbook_path: Path,
    pdf_dir: Path | None,
    extractor: Extractor | None,
    photos_dir: Path | None = None,
    photos_base_url: str = "/images/speakers",
    overwrite: bool = False,
    dry_run: bool = False,
) -> ImportReport:
    workbook = load_workbook(workbook_path)
    ws = workbook[SHEET_SPEAKERS]
    headers = [cell_text(cell.value).rstrip(" *") for cell in ws[1]]
    column_index = {
        column.key: headers.index(column.header) + 1
        for column in SPEAKER_COLUMNS
        if column.header in headers
    }
    report = ImportReport()

    def cell(row: int, key: str):
        return ws.cell(row=row, column=column_index[key])

    for row in range(2, ws.max_row + 1):
        first, last = (
            cell_text(cell(row, "first_name").value),
            cell_text(cell(row, "last_name").value),
        )
        if not first or not last:
            continue
        name = f"{first} {last}"
        speaker_id = slugify(cell_text(cell(row, "id").value)) or slugify(name)
        if not is_yes(cell(row, "consent").value):
            report.skipped.append(f"{name} : pas de consentement")
            continue

        if photos_dir and (overwrite or not cell_text(cell(row, "photo_url").value)):
            photo = find_file(photos_dir, speaker_id, PHOTO_EXTENSIONS)
            if photo:
                cell(row, "photo_url").value = f"{photos_base_url.rstrip('/')}/{photo.name}"
                report.photos.append(name)

        if not pdf_dir or not extractor:
            continue
        targets = [k for k in FIELD_MAP if overwrite or not cell_text(cell(row, k).value)]
        if not targets:
            report.skipped.append(f"{name} : fiche déjà complète")
            continue
        pdf = find_file(pdf_dir, speaker_id, (".pdf",))
        if not pdf:
            report.skipped.append(f"{name} : PDF introuvable ({speaker_id}.pdf)")
            continue
        try:
            extracted = extractor(extract_pdf_text(pdf))
        except Exception as exc:  # erreur réseau, JSON invalide, PDF illisible…
            report.skipped.append(f"{name} : échec de l'extraction ({exc})")
            continue
        for key in targets:
            value = extracted.get(FIELD_MAP[key])
            if isinstance(value, list):
                value = "; ".join(str(item).strip() for item in value if str(item).strip())
            if isinstance(value, str) and value.strip():
                cell(row, key).value = value.strip()
        report.updated.append(name)

    if not dry_run:
        workbook.save(workbook_path)
    return report
